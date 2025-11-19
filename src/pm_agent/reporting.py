"""Documentation/report generation helpers."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import List

from docx import Document
from docx.shared import Inches

from .analysis import ArtifactInfo, CriticalPathResult, MonteCarloResult
from .schema import FocusBoard, ProjectPlan


def _read_file(path: Path, fallback: str) -> str:
    if path.exists():
        return path.read_text(encoding="utf-8")
    return fallback


def generate_documentation_report(
    plan: ProjectPlan,
    focus_board: FocusBoard,
    critical_path: CriticalPathResult,
    monte_carlo: MonteCarloResult,
    artifacts: List[ArtifactInfo],
    output_path: Path,
    project_root: Path,
    gantt_path: Path,
) -> Path:
    """Create a Word report that unifies the project's key artifacts."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading(f"{project_root.name} Project Brief", 0)
    doc.add_paragraph(
        f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')} "
        "via PM Agent automated documentation."
    )

    # Charter
    charter_text = _read_file(
        project_root / "00_project_charter.md",
        "Charter not yet created.",
    )
    doc.add_heading("Project Charter", level=1)
    doc.add_paragraph(charter_text)

    # Plan snapshot
    doc.add_heading("Plan Snapshot", level=1)
    plan_table = doc.add_table(rows=1, cols=5)
    plan_table.style = "Medium List 1 Accent 1"
    plan_table.autofit = True
    hdr_cells = plan_table.rows[0].cells
    hdr_cells[0].text = "ID"
    hdr_cells[1].text = "Title"
    hdr_cells[2].text = "Owner"
    hdr_cells[3].text = "Due"
    hdr_cells[4].text = "Status"

    for task in plan.tasks[:15]:
        row_cells = plan_table.add_row().cells
        row_cells[0].text = str(task.id)
        row_cells[1].text = task.title
        row_cells[2].text = task.assignee or "Unassigned"
        row_cells[3].text = task.due_date.isoformat() if task.due_date else "-"
        row_cells[4].text = task.status

    if len(plan.tasks) > 15:
        doc.add_paragraph(f"(Showing first 15 of {len(plan.tasks)} tasks)")

    # Critical path
    doc.add_heading("Critical Path Analysis", level=1)
    if critical_path.tasks:
        cp_table = doc.add_table(rows=1, cols=5)
        cp_table.style = "Medium Shading 1 Accent 2"
        cp_hdr = cp_table.rows[0].cells
        cp_hdr[0].text = "Task"
        cp_hdr[1].text = "Earliest"
        cp_hdr[2].text = "Latest"
        cp_hdr[3].text = "Slack (d)"
        cp_hdr[4].text = "Critical?"
        for task in critical_path.tasks:
            row = cp_table.add_row().cells
            row[0].text = f"#{task.task_id} {task.title}"
            row[1].text = f"{task.earliest_start} \u2192 {task.earliest_finish}"
            row[2].text = f"{task.latest_start} \u2192 {task.latest_finish}"
            row[3].text = str(task.slack)
            row[4].text = "Yes" if task.is_critical else "No"
    else:
        doc.add_paragraph("No critical path available (plan empty).")

    if critical_path.warnings:
        doc.add_paragraph("Warnings:")
        for warning in critical_path.warnings:
            doc.add_paragraph(f"- {warning}")

    # Monte Carlo risk
    doc.add_heading("Schedule Risk Outlook", level=1)
    risk_paragraph = doc.add_paragraph()
    percentiles = monte_carlo.finish_percentiles
    risk_paragraph.add_run(
        f"P50 completion: {percentiles['p50']}, "
        f"P90 completion: {percentiles['p90']}."
    )
    doc.add_paragraph(
        "Monte Carlo percentiles are calculated from thousands of simulations "
        "using the task-level duration estimates."
    )
    if monte_carlo.risk_drivers:
        doc.add_paragraph("Top risk drivers:")
        for driver in monte_carlo.risk_drivers[:5]:
            doc.add_paragraph(
                f"- #{driver.task_id} {driver.title}: "
                f"correlation {driver.correlation:.2f}, "
                f"avg duration {driver.avg_duration:.1f} days."
            )

    # Focus board
    doc.add_heading("Near-Term Focus", level=1)
    if focus_board.items:
        focus_table = doc.add_table(rows=1, cols=5)
        focus_table.style = "Light List"
        focus_hdr = focus_table.rows[0].cells
        focus_hdr[0].text = "Task"
        focus_hdr[1].text = "Due"
        focus_hdr[2].text = "Owner"
        focus_hdr[3].text = "Status"
        focus_hdr[4].text = "Reason"
        for item in focus_board.items:
            row = focus_table.add_row().cells
            row[0].text = f"#{item.task_id or '?'} {item.title}"
            row[1].text = item.due_date.isoformat() if item.due_date else "-"
            row[2].text = item.assignee or "Unassigned"
            row[3].text = item.status or "-"
            label = item.reason.replace("_", " ").title()
            if item.days_remaining is not None:
                label += f" ({item.days_remaining}d)"
            row[4].text = label
    else:
        doc.add_paragraph("No immediate tasks in the focus list.")

    # Logs and risks
    doc.add_heading("Risks & Issues", level=1)
    doc.add_paragraph(
        _read_file(project_root / "03_risks_issues.md", "No risks/issues captured yet.")
    )

    doc.add_heading("Decision Log Highlights", level=1)
    decisions_text = _read_file(project_root / "04_decisions.log", "No decisions logged.")
    doc.add_paragraph(decisions_text)

    # Artifacts inventory
    doc.add_heading("Artifacts Inventory", level=1)
    if artifacts:
        for art in artifacts:
            doc.add_paragraph(
                f"- {art.path} ({art.size_kb:.1f} KB, updated {art.modified:%Y-%m-%d %H:%M})"
            )
    else:
        doc.add_paragraph("No artifacts stored yet.")

    # Gantt image
    doc.add_heading("Gantt Timeline", level=1)
    if gantt_path.exists():
        doc.add_picture(str(gantt_path), width=Inches(6))
    else:
        doc.add_paragraph("No Gantt chart available.")

    doc.save(output_path)
    return output_path
